# -*- coding: utf-8 -*-
r"""
ArkSwift Outlook Inbox Triage + Domain Tracking
文件名：inbox.py

目录：
C:\code\outreach\
├── inbox.py
├── send.py
├── domains.txt                  # 可选
└── output\
    ├── outreach_emails.xlsx     # 可选，但强烈建议保留
    └── outlook_triage.xlsx      # 本脚本生成

功能：
1) 扫描 Classic Outlook Inbox 全部历史邮件（已读 + 未读）
2) 扫描 Sent Items，建立触达/会话索引
3) 分类：发送失败 / 自动回复 / 真人回复 / 其他
4) 对“高可信真人回复”自动在 Outlook 中打 Flag（无截止日期）
5) 可选读取一列域名，按原顺序生成“域名追踪”子表
6) 域名追踪显示：
   - 是否找到邮箱
   - 是否主动触达
   - 主动触达次数 / 最近触达时间
   - 是否退信 / 失败邮箱
   - 自动回复数量
   - 真人回复数量
   - 最新真人回复
   - 最终结果
   - 历史真人回复（联系人 + 邮件 + 时间 + 正文）
   - 下一步建议
7) 新增“域名统计”：未触达 / 已触达无真人 / 已触达有真人、平均往来量、饼图、柱状图
8) 新增“真人回复历史”：每封真人回复独立一行，避免 Excel 单元格 32767 字符上限
9) 每次运行自动导出 Outlook 中所有已 Flag 邮件：
   - outlook_flagged_replies.jsonl：完整正文，推荐直接交给 GPT
   - outlook_flagged_replies.md：完整正文，人眼/GPT 都方便阅读
   - 主报告新增“Flag邮件”Sheet：结构化索引（正文受 Excel 单元格上限保护）
10) 颜色：
   - 绿色：有真人回复
   - 红色：明确退信 / 发送失败
   - 黄色：已触达但没有真人回复
   - 紫色：没有找到邮箱
   - 蓝色：有邮箱，但 Sent Items 没找到主动触达
"""

import html
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import time

import openpyxl
import pythoncom
import win32com.client
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList


# ============================================================
# 配置
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
REPORT_PATH = OUTPUT_DIR / "outlook_triage.xlsx"
OUTREACH_PATH = OUTPUT_DIR / "outreach_emails.xlsx"
DOMAINS_PATH = BASE_DIR / "domains.txt"
FLAGGED_JSONL_PATH = OUTPUT_DIR / "outlook_flagged_replies.jsonl"
FLAGGED_MD_PATH = OUTPUT_DIR / "outlook_flagged_replies.md"

MAILBOX_ACCOUNT = ""     # 留空 = 默认 Outlook 邮箱
SCAN_INBOX_SUBFOLDERS = True
BODY_PREVIEW_CHARS = 3000
HISTORY_BODY_CHARS = 2500
HISTORY_CELL_MAX_CHARS = 30000   # Excel 单元格硬上限 32767，留安全余量
MAX_INBOX_ITEMS = None   # None = 全部历史
MAX_SENT_ITEMS = None    # None = 全部历史

# 这版会修改 Outlook：只做真人回复 Flag，不删除/移动/标已读/自动回复。
ENABLE_HUMAN_REPLY_FLAG = True

# Outlook 常量
OL_FOLDER_INBOX = 6
OL_FOLDER_SENT_MAIL = 5
OL_MARK_NO_DATE = 4


# ============================================================
# 表头
# ============================================================

TRIAGE_HEADERS = [
    "分类",
    "意向/动作",
    "置信度",
    "收到时间",
    "未读",
    "发件人姓名",
    "发件人邮箱",
    "主题",
    "失败收件邮箱",
    "匹配原因",
    "正文预览",
    "ConversationID",
    "ConversationTopic",
    "MessageClass",
    "Outlook EntryID",
    "所在文件夹",
    "Flag处理",
]

DOMAIN_HEADERS = [
    "域名",
    "原表邮箱",
    "是否找到邮箱",
    "是否主动触达",
    "主动触达次数",
    "最近主动触达时间",
    "主动触达邮箱",
    "是否有退信",
    "失败邮箱",
    "自动回复数",
    "真人回复数",
    "最新真人回复时间",
    "最新联系人",
    "最终结果",
    "历史真人回复",
    "下一步建议",
]


# ============================================================
# 规则
# ============================================================

EMAIL_RE = re.compile(
    r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}",
    re.I,
)

REPLY_PREFIX_RE = re.compile(
    r"^\s*((re|fw|fwd|aw|sv|答复|回复|转发)\s*:\s*)+",
    re.I,
)

BOUNCE_SUBJECT_PATTERNS = [
    r"delivery status notification",
    r"delivery failure",
    r"delivery failed",
    r"undeliverable",
    r"returned mail",
    r"failure notice",
    r"mail delivery failed",
    r"message blocked",
    r"couldn['’]t be delivered",
    r"could not be delivered",
    r"non[- ]delivery",
]

BOUNCE_BODY_PATTERNS = [
    r"couldn['’]t be delivered",
    r"could not be delivered",
    r"wasn['’]t found at",
    r"was not found at",
    r"delivery status notification",
    r"delivery has failed",
    r"delivery failed",
    r"undeliverable",
    r"unknown to address",
    r"recipient address rejected",
    r"user unknown",
    r"mailbox unavailable",
    r"message blocked",
    r"delivery loop",
    r"\b5\.1\.1\b",
    r"\b5\.1\.8\b",
    r"permanent failure",
]

BOUNCE_SENDER_HINTS = [
    "mailer-daemon",
    "mail delivery subsystem",
    "postmaster",
    "microsoft outlook",
    "microsoft exchange",
]

AUTO_SUBJECT_PATTERNS = [
    r"automatic reply",
    r"auto(?:matic)?[- ]?reply",
    r"out of office",
    r"\booo\b",
    r"away from (?:the )?office",
    r"request received",
    r"ticket received",
    r"case received",
    r"support request received",
    r"we(?:'|’)ve received your request",
]

AUTO_BODY_PATTERNS = [
    r"this is an automatic(?:ally generated)? (?:reply|response|message)",
    r"automated (?:reply|response|message)",
    r"we have received your (?:request|message|email)",
    r"we(?:'|’)ve received your (?:request|message|email)",
    r"your (?:request|ticket|case).{0,100}(?:has been|was) received",
    r"we will get back to you",
    r"we(?:'|’)ll get back to you",
    r"we will respond within",
    r"we(?:'|’)ll respond within",
    r"ticket (?:number|#|id)",
    r"support ticket",
    r"do not reply to this automated",
    r"solicitarea dumneavoastr.{0,100}a fost primit",
    r"urmeaz.{0,50}s.{0,20}fie examinat",
    r"vom reveni cu un r.{0,30}spuns",
]

SYSTEM_OTHER_SUBJECT_PATTERNS = [
    r"user activation",
    r"account activation",
    r"activate your account",
    r"verify your email",
    r"email verification",
    r"password reset",
    r"security alert",
    r"sign[- ]in",
    r"login code",
    r"one[- ]time pass",
    r"newsletter",
    r"weekly digest",
    r"daily digest",
]

SYSTEM_OTHER_BODY_PATTERNS = [
    r"new .* account has been created for you",
    r"click .* to activate your account",
    r"select a password",
    r"verify your email",
    r"unsubscribe",
    r"manage your preferences",
]

POSITIVE_PATTERNS = [
    r"\bopen to\b",
    r"\binterested\b",
    r"\bcollaborat",
    r"\bpartnership",
    r"\blink exchange",
    r"\bguest (?:post|contribution)",
    r"\bsponsored",
    r"\bpaid\b",
    r"\bunpaid\b",
    r"\bmedia kit\b",
    r"\beditorial guidelines",
    r"\bpricing\b",
    r"\brate(?:s)?\b",
    r"\bwe accept\b",
    r"\blet me know\b",
]

NEGATIVE_PATTERNS = [
    r"\bnot interested\b",
    r"\bnot open to\b",
    r"\bwe do not accept\b",
    r"\bwe don['’]t accept\b",
    r"\bwe cannot collaborate\b",
    r"\bwe can['’]t collaborate\b",
    r"\bunable to collaborate\b",
    r"\bno partnerships?\b",
    r"\bnot accepting\b",
    r"\bplease remove\b",
    r"\bdo not contact\b",
]

COLLAB_PATTERNS = [
    ("1:1互链", r"\b1\s*[:：]\s*1\b.{0,40}\blink exchange|\blink exchange\b.{0,40}\b1\s*[:：]\s*1\b"),
    ("互链/Link Exchange", r"\blink exchange\b|\breciprocal link\b"),
    ("Guest Post", r"\bguest post\b|\bguest contribution\b"),
    ("Sponsored/付费刊登", r"\bsponsored\b|\bpaid placement\b|\bpaid post\b"),
    ("Article Update/现有文章更新", r"\barticle update\b|\bexisting article\b|\bcontent update\b"),
    ("Affiliate", r"\baffiliate\b"),
    ("Content Exchange", r"\bcontent exchange\b|\bmutual(?:ly)? beneficial content\b"),
    ("Media Kit/Pricing", r"\bmedia kit\b|\bpricing\b|\brates?\b"),
]

PRICE_RE = re.compile(
    r"(?:(?:USD|EUR|GBP)\s*[\$€£]?\s*\d[\d,]*(?:\.\d{1,2})?|"
    r"[\$€£]\s*\d[\d,]*(?:\.\d{1,2})?|"
    r"\d[\d,]*(?:\.\d{1,2})?\s*(?:USD|EUR|GBP))",
    re.I,
)


# ============================================================
# 文本 / 域名工具
# ============================================================

def safe_get(obj, attr, default=""):
    try:
        v = getattr(obj, attr)
        return default if v is None else v
    except Exception:
        return default


def clean_text(v):
    if v is None:
        return ""
    return str(v).replace("\x00", "").strip()


def normalize_subject(subject):
    subject = clean_text(subject)
    subject = REPLY_PREFIX_RE.sub("", subject)
    return re.sub(r"\s+", " ", subject).strip().lower()


def regex_any(patterns, text):
    text = clean_text(text)
    return any(re.search(p, text, re.I | re.S) for p in patterns)


def extract_emails(text):
    if not text:
        return []
    text = str(text).replace(r"\@", "@")
    found = EMAIL_RE.findall(text)
    out, seen = [], set()
    for e in found:
        e = e.lower().strip(" <>[](){}.,;:")
        if e and e not in seen:
            seen.add(e)
            out.append(e)
    return out


def email_domain(email):
    email = clean_text(email).lower()
    if "@" not in email:
        return ""
    return email.rsplit("@", 1)[1].strip(".")


def normalize_domain(value):
    value = clean_text(value).lower()
    value = re.sub(r"^https?://", "", value)
    value = value.split("/", 1)[0]
    value = value.split(":", 1)[0]
    if value.startswith("www."):
        value = value[4:]
    return value.strip(".")


def domain_matches(target, candidate):
    """
    允许：
    - 完全相同
    - target 是 candidate 的子域
    - candidate 是 target 的子域
    """
    t = normalize_domain(target)
    c = normalize_domain(candidate)
    if not t or not c:
        return False
    return t == c or t.endswith("." + c) or c.endswith("." + t)


def fmt_dt(v):
    if v is None:
        return ""
    try:
        return v.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return clean_text(v)


def parse_dt(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
    except Exception:
        return datetime.min


def get_sender_email(item):
    raw = clean_text(safe_get(item, "SenderEmailAddress"))
    if "@" in raw:
        return raw.lower()
    try:
        if clean_text(item.SenderEmailType).upper() == "EX":
            exu = item.Sender.GetExchangeUser()
            if exu:
                smtp = clean_text(exu.PrimarySmtpAddress)
                if smtp:
                    return smtp.lower()
            exl = item.Sender.GetExchangeDistributionList()
            if exl:
                smtp = clean_text(exl.PrimarySmtpAddress)
                if smtp:
                    return smtp.lower()
    except Exception:
        pass
    return raw.lower()


def get_recipient_emails(item):
    """
    尽量解析 Sent Items 中所有 To/Cc 收件人 SMTP 地址。
    """
    out, seen = [], set()

    # 优先 Recipients
    try:
        recipients = item.Recipients
        for i in range(1, recipients.Count + 1):
            try:
                r = recipients.Item(i)
                addr = clean_text(safe_get(r, "Address"))
                if "@" not in addr:
                    try:
                        ae = r.AddressEntry
                        if clean_text(ae.Type).upper() == "EX":
                            exu = ae.GetExchangeUser()
                            if exu:
                                addr = clean_text(exu.PrimarySmtpAddress)
                        else:
                            addr = clean_text(ae.Address)
                    except Exception:
                        pass
                if "@" in addr:
                    addr = addr.lower()
                    if addr not in seen:
                        seen.add(addr)
                        out.append(addr)
            except Exception:
                continue
    except Exception:
        pass

    # fallback: To / CC 字符串
    for field in ("To", "CC"):
        for e in extract_emails(safe_get(item, field, "")):
            if e not in seen:
                seen.add(e)
                out.append(e)

    return out


# ============================================================
# Outlook
# ============================================================

def get_account_and_store(namespace):
    wanted = MAILBOX_ACCOUNT.strip().lower()
    if wanted:
        for account in namespace.Accounts:
            try:
                smtp = clean_text(account.SmtpAddress).lower()
            except Exception:
                smtp = ""
            if smtp == wanted:
                return account, account.DeliveryStore
        raise RuntimeError(f"Outlook 中找不到账号：{MAILBOX_ACCOUNT}")

    inbox = namespace.GetDefaultFolder(OL_FOLDER_INBOX)
    store = inbox.Store
    try:
        account = namespace.Accounts.Item(1)
    except Exception:
        account = None
    return account, store


def walk_folder(folder, recurse=True):
    yield folder
    if not recurse:
        return
    try:
        folders = folder.Folders
        for i in range(1, folders.Count + 1):
            try:
                yield from walk_folder(folders.Item(i), True)
            except Exception:
                continue
    except Exception:
        return


def mark_human_reply_flag(item):
    """
    真正写回 Outlook。
    olMarkNoDate = 4：红旗，无截止日期。
    """
    if not ENABLE_HUMAN_REPLY_FLAG:
        return "未启用"

    try:
        already = bool(safe_get(item, "IsMarkedAsTask", False))
        if already:
            return "已有Flag"

        item.MarkAsTask(OL_MARK_NO_DATE)
        item.Save()
        return "新增Flag"
    except Exception as exc:
        return f"Flag失败: {exc}"


def is_item_flagged(item):
    """
    读取 Outlook 当前 Flag 状态。
    这里不依赖本轮是否刚打 Flag：历史上人工/脚本已经 Flag 的邮件也会被导出。
    """
    try:
        return bool(safe_get(item, "IsMarkedAsTask", False))
    except Exception:
        return False


# ============================================================
# Outreach XLSX 映射
# ============================================================

def find_header(ws, aliases):
    norm_aliases = {
        re.sub(r"\s+", "", str(x)).lower()
        for x in aliases
    }
    for cell in ws[1]:
        v = re.sub(r"\s+", "", clean_text(cell.value)).lower()
        if v in norm_aliases:
            return cell.column
    return None


def load_outreach_mapping():
    """
    domain -> [emails...]
    同时建立 email -> domains
    """
    domain_to_emails = defaultdict(list)
    email_to_domains = defaultdict(set)

    if not OUTREACH_PATH.exists():
        return domain_to_emails, email_to_domains

    try:
        wb = openpyxl.load_workbook(OUTREACH_PATH, read_only=True, data_only=True)
        ws = wb.active
    except Exception as exc:
        print(f"[警告] 无法读取 outreach_emails.xlsx：{exc}")
        return domain_to_emails, email_to_domains

    domain_col = find_header(ws, ["Domain (域名)", "Domain", "域名"])
    to_col = find_header(ws, ["To (收件邮箱)", "To", "收件邮箱", "Email", "邮箱"])

    if not domain_col or not to_col:
        print("[警告] outreach_emails.xlsx 没找到 Domain/To 列。")
        return domain_to_emails, email_to_domains

    for row in range(2, ws.max_row + 1):
        d = normalize_domain(ws.cell(row, domain_col).value)
        if not d:
            continue
        emails = extract_emails(ws.cell(row, to_col).value)
        for e in emails:
            if e not in domain_to_emails[d]:
                domain_to_emails[d].append(e)
            email_to_domains[e].add(d)

    return domain_to_emails, email_to_domains


# ============================================================
# Sent Items 索引
# ============================================================

def build_sent_index(sent_folder):
    sent_records = []
    conv_ids = set()
    subjects = set()
    recipient_to_sent = defaultdict(list)
    recipient_domain_to_sent = defaultdict(list)
    conv_to_sent = defaultdict(list)
    subject_to_sent = defaultdict(list)

    print("正在建立 Sent Items 会话索引...")

    items = sent_folder.Items
    try:
        items.Sort("[SentOn]", True)
    except Exception:
        pass

    total = items.Count
    limit = total if MAX_SENT_ITEMS is None else min(total, MAX_SENT_ITEMS)

    for i in range(1, limit + 1):
        try:
            item = items.Item(i)
        except Exception:
            continue

        mc = clean_text(safe_get(item, "MessageClass"))
        if not mc.startswith("IPM.Note"):
            continue

        subject = clean_text(safe_get(item, "Subject"))
        conv_id = clean_text(safe_get(item, "ConversationID"))
        conv_topic = clean_text(safe_get(item, "ConversationTopic"))
        sent_time = fmt_dt(safe_get(item, "SentOn", None))
        recipients = get_recipient_emails(item)

        record = {
            "subject": subject,
            "norm_subject": normalize_subject(subject),
            "conv_id": conv_id,
            "conv_topic": conv_topic,
            "norm_conv_topic": normalize_subject(conv_topic),
            "sent_time": sent_time,
            "recipients": recipients,
            "entry_id": clean_text(safe_get(item, "EntryID")),
        }

        sent_records.append(record)

        if conv_id:
            conv_ids.add(conv_id)
            conv_to_sent[conv_id].append(record)

        for s in (record["norm_subject"], record["norm_conv_topic"]):
            if s:
                subjects.add(s)
                subject_to_sent[s].append(record)

        for e in recipients:
            recipient_to_sent[e].append(record)
            d = email_domain(e)
            if d:
                recipient_domain_to_sent[d].append(record)

        if i % 500 == 0:
            print(f"  已索引 {i}/{limit}")

    print(
        f"Sent Items：{len(sent_records)} 封；"
        f"{len(conv_ids)} 个会话；{len(subjects)} 个主题。\n"
    )

    return {
        "records": sent_records,
        "conv_ids": conv_ids,
        "subjects": subjects,
        "recipient_to_sent": recipient_to_sent,
        "recipient_domain_to_sent": recipient_domain_to_sent,
        "conv_to_sent": conv_to_sent,
        "subject_to_sent": subject_to_sent,
    }


# ============================================================
# 分类
# ============================================================

def has_reply_prefix(subject):
    return bool(REPLY_PREFIX_RE.match(clean_text(subject)))


def _sent_record_key(rec):
    return rec.get("entry_id") or (
        rec.get("sent_time", ""),
        rec.get("subject", ""),
        tuple(rec.get("recipients", [])),
    )


def _recipient_emails(records):
    out = []
    seen = set()
    for rec in records:
        for e in rec.get("recipients", []):
            e = clean_text(e).lower()
            if e and e not in seen:
                seen.add(e)
                out.append(e)
    return out


def _sender_compatible_with_records(sender_email, records):
    """发件人必须与这些 Sent 记录中的某个实际收件人邮箱/域名相符。"""
    sender_email = clean_text(sender_email).lower()
    if not sender_email:
        return False

    sender_domain = email_domain(sender_email)
    for e in _recipient_emails(records):
        if sender_email == e:
            return True
        rd = email_domain(e)
        if sender_domain and rd and domain_matches(sender_domain, rd):
            return True
    return False


def _conversation_is_narrow(records):
    """
    ConversationID 只有在它没有横跨大量不同收件人时才可单独辅助判断。
    防止批量外联使用相同主题/会话字段时串线。
    """
    recipients = _recipient_emails(records)
    recipient_domains = {email_domain(e) for e in recipients if email_domain(e)}
    return bool(records) and len(recipients) <= 3 and len(recipient_domains) <= 1


def extract_failed_recipients(subject, body, own_email=""):
    """
    只从 NDR/退信的“失败收件人上下文”提取邮箱。
    不再 fallback 为正文里随便出现的第一个邮箱，避免把引用正文中的联系人错判为失败邮箱。
    """
    combined = f"{subject}\n{body}".replace(r"\@", "@")
    own_email = clean_text(own_email).lower()

    patterns = [
        r"your message to\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"message to\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"recipient(?: address)?\s*[:=]\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"delivery to\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"couldn['’]t be delivered to\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"could not be delivered to\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"final-recipient\s*:\s*rfc822\s*;\s*([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"original-recipient\s*:\s*rfc822\s*;\s*([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
        r"failed recipient\s*[:=]\s*[<\[]?([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})",
    ]

    out = []
    seen = set()
    for p in patterns:
        for m in re.finditer(p, combined, re.I | re.S):
            e = m.group(1).lower().strip(" <>[](){}.,;:")
            if not e or e == own_email:
                continue
            if any(x in e for x in ["mailer-daemon", "postmaster", "microsoftexchange"]):
                continue
            if e not in seen:
                seen.add(e)
                out.append(e)

    # 兜底：只看包含明确失败关键词的局部行，而不是整封邮件任意邮箱。
    if not out:
        lines = combined.splitlines()
        failure_hint = re.compile(
            r"undeliver|delivery has failed|delivery failed|couldn['’]t be delivered|could not be delivered|"
            r"recipient address rejected|user unknown|mailbox unavailable|5\.1\.1|5\.1\.8|failed recipient",
            re.I,
        )
        for idx, line in enumerate(lines):
            if not failure_hint.search(line):
                continue
            context = "\n".join(lines[max(0, idx - 1): min(len(lines), idx + 2)])
            for e in extract_emails(context):
                if e == own_email:
                    continue
                if any(x in e for x in ["mailer-daemon", "postmaster", "microsoftexchange"]):
                    continue
                if e not in seen:
                    seen.add(e)
                    out.append(e)

    return out


def classify_item(item, sent_index, own_email):
    subject = clean_text(safe_get(item, "Subject"))
    body = clean_text(safe_get(item, "Body"))
    sender_name = clean_text(safe_get(item, "SenderName"))
    sender_email = get_sender_email(item)
    mc = clean_text(safe_get(item, "MessageClass"))
    conv_id = clean_text(safe_get(item, "ConversationID"))
    conv_topic = clean_text(safe_get(item, "ConversationTopic"))
    combined = f"{subject}\n{sender_name}\n{sender_email}\n{body}"

    # 1) 退信
    is_ndr = mc.upper().startswith("REPORT.") and "NDR" in mc.upper()
    bounce_sender = any(
        hint in f"{sender_name} {sender_email}".lower()
        for hint in BOUNCE_SENDER_HINTS
    )
    bounce_subject = regex_any(BOUNCE_SUBJECT_PATTERNS, subject)
    bounce_body = regex_any(BOUNCE_BODY_PATTERNS, body)

    if is_ndr or bounce_subject or (bounce_sender and bounce_body):
        failed_recipients = extract_failed_recipients(subject, body, own_email)
        return {
            "category": "发送失败",
            "confidence": "高",
            "reason": "NDR/退信发件人/典型投递失败文本命中。",
            "failed_recipient": "; ".join(failed_recipients),
            "failed_recipients": failed_recipients,
            "intent": "需要重新找邮箱",
            "is_human": False,
        }

    # 2) 自动回复
    auto_subject = regex_any(AUTO_SUBJECT_PATTERNS, subject)
    auto_body = regex_any(AUTO_BODY_PATTERNS, body)

    if auto_subject or auto_body:
        return {
            "category": "自动回复",
            "confidence": "高" if auto_subject and auto_body else "中高",
            "reason": "命中自动回复/工单回执/OOO 模板。",
            "failed_recipient": "",
            "failed_recipients": [],
            "intent": "无需处理",
            "is_human": False,
        }

    # 3) 系统其他
    if (
        regex_any(SYSTEM_OTHER_SUBJECT_PATTERNS, subject)
        or regex_any(SYSTEM_OTHER_BODY_PATTERNS, body)
    ):
        return {
            "category": "其他",
            "confidence": "高",
            "reason": "账户激活/验证/Newsletter 等系统通知。",
            "failed_recipient": "",
            "failed_recipients": [],
            "intent": "系统/其他",
            "is_human": False,
        }

    # 4) 真人回复：禁止“Subject 相同 = 真人回复”的宽松逻辑
    norm_subject = normalize_subject(subject)
    norm_topic = normalize_subject(conv_topic)

    conv_records = sent_index["conv_to_sent"].get(conv_id, []) if conv_id else []
    subject_records = []
    seen_sent = set()
    for key in (norm_subject, norm_topic):
        if not key:
            continue
        for rec in sent_index["subject_to_sent"].get(key, []):
            rk = _sent_record_key(rec)
            if rk not in seen_sent:
                seen_sent.add(rk)
                subject_records.append(rec)

    sender_exact_records = sent_index["recipient_to_sent"].get(sender_email, []) if sender_email else []
    sender_was_exact_recipient = bool(sender_exact_records)
    sender_matches_conv = _sender_compatible_with_records(sender_email, conv_records)
    sender_matches_subject = _sender_compatible_with_records(sender_email, subject_records)

    # ConversationID 可作为强证据，但仅限窄会话；否则必须由 sender 邮箱/域名交叉验证。
    narrow_conv_reply = bool(
        conv_records
        and _conversation_is_narrow(conv_records)
        and has_reply_prefix(subject)
    )

    strong_reply = (
        sender_was_exact_recipient
        or sender_matches_conv
        or sender_matches_subject
        or narrow_conv_reply
    )

    if strong_reply:
        positives = sum(
            bool(re.search(p, combined, re.I | re.S))
            for p in POSITIVE_PATTERNS
        )
        negatives = sum(
            bool(re.search(p, combined, re.I | re.S))
            for p in NEGATIVE_PATTERNS
        )

        if negatives:
            intent = "明确拒绝候选"
        elif positives >= 2:
            intent = "高合作意向候选"
        else:
            intent = "真人回复-待判断"

        if sender_was_exact_recipient:
            reason = "发件邮箱精确等于历史主动触达收件人。"
            conf = "高"
        elif sender_matches_conv:
            reason = "ConversationID 命中，且发件邮箱域名与该会话实际收件人匹配。"
            conf = "高"
        elif sender_matches_subject:
            reason = "主题命中 Sent，且发件邮箱域名与该主题实际收件人匹配。"
            conf = "中高"
        else:
            reason = "ConversationID 为单一/窄目标会话，且邮件具有回复前缀。"
            conf = "中高"

        return {
            "category": "真人回复",
            "confidence": conf,
            "reason": reason,
            "failed_recipient": "",
            "failed_recipients": [],
            "intent": intent,
            "is_human": True,
        }

    # 5) 其他
    return {
        "category": "其他",
        "confidence": "中",
        "reason": "未命中退信/自动回复；同时缺少‘发件人 ↔ 已发送收件人’交叉证据，未判为真人回复。",
        "failed_recipient": "",
        "failed_recipients": [],
        "intent": "其他/待抽查",
        "is_human": False,
    }


# ============================================================
# 合作结果提取
# ============================================================

def infer_latest_outcome(record):
    body = record.get("body", "")
    subject = record.get("subject", "")
    combined = f"{subject}\n{body}"

    # 明确拒绝优先
    if regex_any(NEGATIVE_PATTERNS, combined):
        label = "明确拒绝"
    else:
        labels = []
        for name, p in COLLAB_PATTERNS:
            if re.search(p, combined, re.I | re.S):
                labels.append(name)

        prices = PRICE_RE.findall(combined)

        if labels:
            label = "；".join(dict.fromkeys(labels))
        else:
            label = "真人回复，合作方式待人工判断"

        if prices:
            unique_prices = list(dict.fromkeys([clean_text(x) for x in prices]))
            label += "；报价/金额：" + " / ".join(unique_prices[:5])

    return label


# ============================================================
# 域名关联
# ============================================================

def build_domain_aliases(domain, domain_to_emails):
    """
    域名归属采用“严格模式”：
    - 原表邮箱：只取这个 Domain 的精确映射；允许 mapped_domain 是 target 的真正子域。
    - 不再把父域/公共邮箱提供商域名（gmail.com 等）扩散到其它目标。
    """
    d = normalize_domain(domain)
    emails = []

    for mapped_domain, mapped_emails in domain_to_emails.items():
        md = normalize_domain(mapped_domain)
        # 精确匹配，或映射项确实是当前网站的子域。
        if md == d or (d and md.endswith("." + d)):
            for e in mapped_emails:
                e = clean_text(e).lower()
                if e and e not in emails:
                    emails.append(e)

    # 只有与网站域本身相符的邮箱域名，才允许做“域名级”匹配。
    trusted_email_domains = []
    for e in emails:
        ed = email_domain(e)
        if ed and domain_matches(d, ed) and ed not in trusted_email_domains:
            trusted_email_domains.append(ed)

    return d, emails, trusted_email_domains


def _domain_key_matches_target(target_domain, candidate_domain):
    """candidate_domain 必须等于 target，或是 target 的子域；不反向吞父域。"""
    t = normalize_domain(target_domain)
    c = normalize_domain(candidate_domain)
    return bool(t and c and (c == t or c.endswith("." + t)))


def _email_belongs_to_domain(email, domain, known_emails):
    email = clean_text(email).lower()
    if not email:
        return False
    if email in known_emails:
        return True
    ed = email_domain(email)
    return bool(ed and _domain_key_matches_target(domain, ed))


def _sent_record_belongs_to_domain(rec, domain, known_emails):
    return any(
        _email_belongs_to_domain(e, domain, known_emails)
        for e in rec.get("recipients", [])
    )


def sent_matches_domain(domain, known_emails, known_email_domains, sent_index):
    """
    只允许两种触达归属：
    1) 收件邮箱精确等于 outreach 表中的该域名联系人；
    2) 收件邮箱域名与目标网站域名相同/子域关系。

    known_email_domains 参数保留兼容，但不再用来把 gmail.com / group.one 等第三方域名扩大匹配。
    """
    matches = []
    seen = set()
    d = normalize_domain(domain)

    # 1) 最可靠：原表实际邮箱精确匹配
    for e in known_emails:
        for rec in sent_index["recipient_to_sent"].get(e, []):
            key = _sent_record_key(rec)
            if key not in seen:
                seen.add(key)
                matches.append(rec)

    # 2) 网站域名匹配。直接走 recipient_domain_to_sent 索引；
    #    1000 个域名时避免每个域名都全扫一遍 Sent Items。
    for recipient_domain, records in sent_index["recipient_domain_to_sent"].items():
        if not _domain_key_matches_target(d, recipient_domain):
            continue
        for rec in records:
            key = _sent_record_key(rec)
            if key not in seen:
                seen.add(key)
                matches.append(rec)

    return sorted(matches, key=lambda r: parse_dt(r["sent_time"]))


def _inbox_record_key(record):
    return record.get("entry_id") or (
        record.get("received_time", ""),
        record.get("sender_email", ""),
        record.get("subject", ""),
    )


def build_inbox_match_index(inbox_records):
    """为 1000+ 域名追踪建立内存索引，避免每个域名都全扫所有 Inbox。"""
    idx = {
        "sender_email": defaultdict(list),
        "sender_domain": defaultdict(list),
        "failed_email": defaultdict(list),
        "failed_domain": defaultdict(list),
        "conversation_id": defaultdict(list),
    }

    for r in inbox_records:
        sender = clean_text(r.get("sender_email", "")).lower()
        if sender:
            idx["sender_email"][sender].append(r)
            sd = email_domain(sender)
            if sd:
                idx["sender_domain"][sd].append(r)

        for failed in (r.get("failed_recipients") or extract_emails(r.get("failed_recipient", ""))):
            failed = clean_text(failed).lower()
            if not failed:
                continue
            idx["failed_email"][failed].append(r)
            fd = email_domain(failed)
            if fd:
                idx["failed_domain"][fd].append(r)

        conv_id = clean_text(r.get("conversation_id", ""))
        if conv_id:
            idx["conversation_id"][conv_id].append(r)

    return idx


def candidate_inbox_records_for_domain(domain, known_emails, sent_matches, inbox_match_index):
    """先用邮箱/域名/ConversationID 索引缩小候选，再交给严格匹配函数复核。"""
    d = normalize_domain(domain)
    out = []
    seen = set()

    def add(records):
        for r in records:
            k = _inbox_record_key(r)
            if k not in seen:
                seen.add(k)
                out.append(r)

    for e in known_emails:
        add(inbox_match_index["sender_email"].get(e, []))
        add(inbox_match_index["failed_email"].get(e, []))

    for sender_domain, records in inbox_match_index["sender_domain"].items():
        if _domain_key_matches_target(d, sender_domain):
            add(records)

    for failed_domain, records in inbox_match_index["failed_domain"].items():
        if _domain_key_matches_target(d, failed_domain):
            add(records)

    for conv_id in {x.get("conv_id") for x in sent_matches if x.get("conv_id")}:
        add(inbox_match_index["conversation_id"].get(conv_id, []))

    return out


def _conversation_is_exclusive_to_domain(conv_id, domain, known_emails, sent_matches, sent_index):
    """
    ConversationID 只有在该会话的所有 Sent 记录都属于当前域名时，才允许作为跨邮箱回复的归属证据。
    这样即使 Outlook 把同主题批量邮件放进同一 Conversation，也不会串域名。
    """
    if not conv_id:
        return False

    conv_records = sent_index["conv_to_sent"].get(conv_id, [])
    if not conv_records:
        return False

    target_keys = {_sent_record_key(x) for x in sent_matches}
    conv_keys = {_sent_record_key(x) for x in conv_records}
    if not conv_keys.issubset(target_keys):
        return False

    return all(
        _sent_record_belongs_to_domain(rec, domain, known_emails)
        for rec in conv_records
    )


def inbox_record_domain_match_reason(
    record,
    domain,
    known_emails,
    known_email_domains,
    sent_matches,
    sent_index,
):
    """
    返回匹配依据；空字符串表示不属于该域名。

    重要：彻底取消 Subject-only 的域名归属。
    """
    d = normalize_domain(domain)
    sender = clean_text(record.get("sender_email", "")).lower()
    failed_emails = record.get("failed_recipients") or extract_emails(record.get("failed_recipient", ""))

    # A. 精确邮箱优先
    if sender and sender in known_emails:
        return "发件邮箱 = 该域名原表联系人邮箱"

    for failed in failed_emails:
        if failed in known_emails:
            return "失败收件邮箱 = 该域名原表联系人邮箱"

    # B. 网站域名正则匹配
    if sender and _email_belongs_to_domain(sender, d, []):
        return "发件邮箱域名匹配目标网站域名"

    for failed in failed_emails:
        if _email_belongs_to_domain(failed, d, []):
            return "失败收件邮箱域名匹配目标网站域名"

    # C. 跨邮箱/工单系统回复：ConversationID 必须对当前域名独占
    conv_id = clean_text(record.get("conversation_id", ""))
    if _conversation_is_exclusive_to_domain(
        conv_id,
        d,
        known_emails,
        sent_matches,
        sent_index,
    ):
        return "ConversationID 仅对应当前域名的 Sent 会话"

    return ""


def inbox_record_matches_domain(
    record,
    domain,
    known_emails,
    known_email_domains,
    sent_matches,
    sent_index=None,
):
    """兼容旧调用：返回 bool。"""
    if sent_index is None:
        return False
    return bool(inbox_record_domain_match_reason(
        record,
        domain,
        known_emails,
        known_email_domains,
        sent_matches,
        sent_index,
    ))


# ============================================================
# Excel
# ============================================================

def make_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for name in ["汇总", "全部", "发送失败", "自动回复", "真人回复", "其他"]:
        wb.create_sheet(name)

    for name in ["全部", "发送失败", "自动回复", "真人回复", "其他"]:
        ws = wb[name]
        ws.append(TRIAGE_HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(TRIAGE_HEADERS))}1"

    return wb


def triage_record_row(r):
    return [
        r["category"],
        r["intent"],
        r["confidence"],
        r["received_time"],
        r["unread"],
        r["sender_name"],
        r["sender_email"],
        r["subject"],
        r["failed_recipient"],
        r["reason"],
        r["body_preview"],
        r["conversation_id"],
        r["conversation_topic"],
        r["message_class"],
        r["entry_id"],
        r["folder"],
        r["flag_result"],
    ]


def apply_triage_widths(ws):
    widths = {
        1: 12, 2: 20, 3: 10, 4: 20, 5: 8, 6: 24, 7: 32,
        8: 52, 9: 32, 10: 52, 11: 90, 12: 32, 13: 42,
        14: 28, 15: 34, 16: 30, 17: 16,
    }
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def write_summary(wb, counts, unread, flag_new, flag_existing, flag_failed):
    ws = wb["汇总"]
    ws["A1"] = "ArkSwift Outlook 邮件整理"
    ws["A1"].font = Font(bold=True, size=16)

    rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("扫描邮件总数", sum(counts.values())),
        ("其中未读", unread),
        ("发送失败", counts["发送失败"]),
        ("自动回复", counts["自动回复"]),
        ("真人回复", counts["真人回复"]),
        ("其他", counts["其他"]),
        ("新增真人 Flag", flag_new),
        ("已有 Flag", flag_existing),
        ("Flag 失败", flag_failed),
    ]

    for idx, (k, v) in enumerate(rows, start=3):
        ws.cell(idx, 1).value = k
        ws.cell(idx, 2).value = v

    ws["A15"] = "说明"
    ws["A15"].font = Font(bold=True)
    ws["A16"] = (
        "脚本不会删除、移动、标已读或自动回复邮件。"
        "唯一写回 Outlook 的动作是真人回复 Flag。"
    )
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80


GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
PURPLE = PatternFill("solid", fgColor="E4DFEC")
BLUE = PatternFill("solid", fgColor="DDEBF7")


def make_history_text(human_records):
    """
    域名追踪单元格给出可读摘要；详细全文另写“真人回复历史”Sheet。
    严格控制在 Excel 32767 字符上限以内，避免历史列异常/不可读。
    """
    parts = []
    ordered = sorted(human_records, key=lambda x: parse_dt(x["received_time"]))

    for idx, r in enumerate(ordered, start=1):
        part = (
            f"===== 真人回复 {idx}/{len(ordered)} =====\n"
            f"时间: {r['received_time']}\n"
            f"联系人: {r['sender_name']}\n"
            f"邮箱: {r['sender_email']}\n"
            f"主题: {r['subject']}\n"
            f"意向预判: {r['intent']}\n"
            f"域名匹配: {r.get('_domain_match_reason', '')}\n"
            f"Flag: {r['flag_result']}\n\n"
            f"{r['body'][:HISTORY_BODY_CHARS]}"
        )
        candidate = "\n\n".join(parts + [part])
        if len(candidate) > HISTORY_CELL_MAX_CHARS:
            remaining = len(ordered) - len(parts)
            tail = f"\n\n[单元格长度已达安全上限；其余 {remaining} 封详见『真人回复历史』Sheet]"
            if len("\n\n".join(parts) + tail) <= HISTORY_CELL_MAX_CHARS:
                parts.append(tail.strip())
            break
        parts.append(part)

    return "\n\n".join(parts)


def _safe_pct(n, d):
    return n / d if d else 0


def _domain_status(touched, humans):
    if not touched:
        return "未主动触达"
    if humans:
        return "已触达且有真人回复"
    return "已触达但无真人回复"


def write_domain_statistics(wb, domain_stats):
    if "域名统计" in wb.sheetnames:
        del wb["域名统计"]

    ws = wb.create_sheet("域名统计")
    ws["A1"] = "ArkSwift 域名触达与回复统计"
    ws["A1"].font = Font(bold=True, size=16)

    total_domains = len(domain_stats)
    untouched = sum(1 for x in domain_stats if not x["touched"])
    touched_no_human = sum(1 for x in domain_stats if x["touched"] and x["human_count"] == 0)
    touched_human = sum(1 for x in domain_stats if x["touched"] and x["human_count"] > 0)
    touched_domains = touched_no_human + touched_human
    known_email_not_touched = sum(1 for x in domain_stats if (not x["touched"]) and x["has_known_email"])
    no_email_not_touched = sum(1 for x in domain_stats if (not x["touched"]) and (not x["has_known_email"]))
    bounce_domains = sum(1 for x in domain_stats if x["bounce_count"] > 0)

    sent_total = sum(x["sent_count"] for x in domain_stats)
    auto_total = sum(x["auto_count"] for x in domain_stats)
    human_total = sum(x["human_count"] for x in domain_stats)
    bounce_total = sum(x["bounce_count"] for x in domain_stats)
    exchange_total = sent_total + auto_total + human_total
    effective_exchange_total = sent_total + human_total

    kpis = [
        ("总域名数", total_domains),
        ("未主动触达", untouched),
        ("已触达但无真人回复", touched_no_human),
        ("已触达且有真人回复", touched_human),
        ("触达率", _safe_pct(touched_domains, total_domains)),
        ("真人回复域名率（占已触达）", _safe_pct(touched_human, touched_domains)),
        ("有邮箱但未触达", known_email_not_touched),
        ("无邮箱且未触达", no_email_not_touched),
        ("出现退信的域名数", bounce_domains),
        ("主动发送总数", sent_total),
        ("自动回复总数", auto_total),
        ("真人回复总数", human_total),
        ("退信记录总数", bounce_total),
        ("平均主动发送/已触达域名", sent_total / touched_domains if touched_domains else 0),
        ("平均真人回复/有真人回复域名", human_total / touched_human if touched_human else 0),
        ("平均往来邮件/全部域名", exchange_total / total_domains if total_domains else 0),
        ("平均往来邮件/已触达域名", exchange_total / touched_domains if touched_domains else 0),
        ("平均有效往来/已触达域名（发送+真人）", effective_exchange_total / touched_domains if touched_domains else 0),
    ]

    ws["A3"] = "核心指标"
    ws["A3"].font = Font(bold=True)
    for idx, (k, v) in enumerate(kpis, start=4):
        ws.cell(idx, 1).value = k
        ws.cell(idx, 2).value = v
        if "率" in k:
            ws.cell(idx, 2).number_format = "0.0%"
        elif "平均" in k:
            ws.cell(idx, 2).number_format = "0.00"

    dist_start = 4
    ws.cell(dist_start, 4).value = "状态"
    ws.cell(dist_start, 5).value = "域名数"
    ws.cell(dist_start, 6).value = "占比"
    dist = [
        ("未主动触达", untouched),
        ("已触达但无真人回复", touched_no_human),
        ("已触达且有真人回复", touched_human),
    ]
    for r, (label, count) in enumerate(dist, start=dist_start + 1):
        ws.cell(r, 4).value = label
        ws.cell(r, 5).value = count
        ws.cell(r, 6).value = _safe_pct(count, total_domains)
        ws.cell(r, 6).number_format = "0.0%"

    # 饼图：1000 个域名处于哪一种状态
    pie = PieChart()
    pie.title = "域名触达状态分布"
    pie.height = 8
    pie.width = 11
    labels = Reference(ws, min_col=4, min_row=dist_start + 1, max_row=dist_start + len(dist))
    data = Reference(ws, min_col=5, min_row=dist_start, max_row=dist_start + len(dist))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = True
    ws.add_chart(pie, "H3")

    # 柱状图：同一组状态便于直接比较数量
    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "域名状态数量对比"
    bar.y_axis.title = "域名数"
    bar.x_axis.title = "状态"
    bar.height = 8
    bar.width = 13
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(labels)
    ws.add_chart(bar, "H20")

    # Top 20：按总往来量排序，便于找“聊得最多”的域名
    top = sorted(
        domain_stats,
        key=lambda x: (x["sent_count"] + x["auto_count"] + x["human_count"], x["human_count"]),
        reverse=True,
    )[:20]
    top_start = 23
    headers = ["域名", "状态", "主动发送", "自动回复", "真人回复", "退信", "往来邮件"]
    for c, h in enumerate(headers, start=1):
        ws.cell(top_start, c).value = h
    for r, x in enumerate(top, start=top_start + 1):
        ws.cell(r, 1).value = x["domain"]
        ws.cell(r, 2).value = x["status"]
        ws.cell(r, 3).value = x["sent_count"]
        ws.cell(r, 4).value = x["auto_count"]
        ws.cell(r, 5).value = x["human_count"]
        ws.cell(r, 6).value = x["bounce_count"]
        ws.cell(r, 7).value = x["sent_count"] + x["auto_count"] + x["human_count"]

    # 样式
    for cell in list(ws[3]) + list(ws[dist_start]) + list(ws[top_start]):
        if cell.value is not None:
            cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14


def write_human_history_sheet(wb, history_rows):
    if "真人回复历史" in wb.sheetnames:
        del wb["真人回复历史"]

    ws = wb.create_sheet("真人回复历史")
    headers = [
        "域名", "收到时间", "联系人", "发件邮箱", "主题", "意向预判",
        "域名匹配依据", "Flag处理", "正文",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    for x in history_rows:
        r = x["record"]
        ws.append([
            x["domain"],
            r["received_time"],
            r["sender_name"],
            r["sender_email"],
            r["subject"],
            r["intent"],
            x["reason"],
            r["flag_result"],
            r["body"][:30000],
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    widths = [34, 22, 28, 36, 55, 24, 46, 18, 100]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)




def _flagged_sort_key(record):
    return parse_dt(record.get("received_time", ""))


def write_flagged_mail_sheet(wb, flagged_records):
    """
    主报告中的结构化浏览表。
    注意：Excel 单元格最大 32767 字符，因此正文列仅保留最多 30000 字符；
    完整正文一定写入 JSONL / Markdown，不在这里丢失。
    """
    if "Flag邮件" in wb.sheetnames:
        del wb["Flag邮件"]

    ws = wb.create_sheet("Flag邮件")
    headers = [
        "序号",
        "收到时间",
        "发件人姓名",
        "发件人邮箱",
        "发件邮箱域名",
        "主题",
        "当前分类",
        "意向/动作",
        "置信度",
        "未读",
        "所在文件夹",
        "ConversationID",
        "ConversationTopic",
        "Outlook EntryID",
        "正文字符数",
        "正文（Excel安全截断）",
        "完整正文文件",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    ordered = sorted(flagged_records, key=_flagged_sort_key, reverse=True)
    for idx, r in enumerate(ordered, start=1):
        body = r.get("body", "") or ""
        ws.append([
            idx,
            r.get("received_time", ""),
            r.get("sender_name", ""),
            r.get("sender_email", ""),
            email_domain(r.get("sender_email", "")),
            r.get("subject", ""),
            r.get("category", ""),
            r.get("intent", ""),
            r.get("confidence", ""),
            r.get("unread", ""),
            r.get("folder", ""),
            r.get("conversation_id", ""),
            r.get("conversation_topic", ""),
            r.get("entry_id", ""),
            len(body),
            body[:30000],
            "outlook_flagged_replies.jsonl / outlook_flagged_replies.md",
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    widths = [8, 22, 28, 36, 28, 60, 14, 24, 12, 8, 34, 34, 44, 34, 14, 100, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def export_flagged_mail_files(flagged_records):
    """
    导出所有当前处于 Flag 状态的 Inbox 邮件。

    JSONL 是“机器/GPT 主数据源”：每行一封邮件，保留完整纯文本正文。
    Markdown 是“人眼/GPT 可读版”：按最新时间在前，完整正文不截断。

    这里故意只做结构化整理，不尝试判断合作模式、报价、是否值得合作等业务结论。
    """
    ordered = sorted(flagged_records, key=_flagged_sort_key, reverse=True)
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # JSONL：一封邮件一行，正文完整保留。
    with FLAGGED_JSONL_PATH.open("w", encoding="utf-8") as f:
        for idx, r in enumerate(ordered, start=1):
            body = r.get("body", "") or ""
            obj = {
                "record_no": idx,
                "flagged": True,
                "received_time": r.get("received_time", ""),
                "sender_name": r.get("sender_name", ""),
                "sender_email": r.get("sender_email", ""),
                "sender_domain": email_domain(r.get("sender_email", "")),
                "subject": r.get("subject", ""),
                "classification": r.get("category", ""),
                "intent": r.get("intent", ""),
                "confidence": r.get("confidence", ""),
                "unread": r.get("unread", ""),
                "folder": r.get("folder", ""),
                "conversation_id": r.get("conversation_id", ""),
                "conversation_topic": r.get("conversation_topic", ""),
                "outlook_entry_id": r.get("entry_id", ""),
                "body_chars": len(body),
                "body": body,
            }
            f.write(json.dumps(obj, ensure_ascii=False) + "\n")

    # Markdown：完整正文，方便直接拖给 GPT 阅读。
    with FLAGGED_MD_PATH.open("w", encoding="utf-8") as f:
        f.write("# Outlook Flag 邮件完整导出\n\n")
        f.write(f"- 生成时间：{generated_at}\n")
        f.write(f"- Flag 邮件数：{len(ordered)}\n")
        f.write("- 说明：本文件仅整理 Outlook 中当前已经 Flag 的邮件，不做合作模式/报价判断。\n\n")

        for idx, r in enumerate(ordered, start=1):
            body = r.get("body", "") or ""
            f.write(f"## Flag 邮件 {idx}/{len(ordered)}\n\n")
            f.write(f"- 收到时间：{r.get('received_time', '')}\n")
            f.write(f"- 联系人：{r.get('sender_name', '')}\n")
            f.write(f"- 发件邮箱：{r.get('sender_email', '')}\n")
            f.write(f"- 发件邮箱域名：{email_domain(r.get('sender_email', ''))}\n")
            f.write(f"- 主题：{r.get('subject', '')}\n")
            f.write(f"- 当前分类：{r.get('category', '')}\n")
            f.write(f"- 意向/动作：{r.get('intent', '')}\n")
            f.write(f"- 置信度：{r.get('confidence', '')}\n")
            f.write(f"- 未读：{r.get('unread', '')}\n")
            f.write(f"- 文件夹：{r.get('folder', '')}\n")
            f.write(f"- ConversationID：{r.get('conversation_id', '')}\n")
            f.write(f"- ConversationTopic：{r.get('conversation_topic', '')}\n")
            f.write(f"- Outlook EntryID：{r.get('entry_id', '')}\n")
            f.write(f"- 正文字符数：{len(body)}\n\n")
            f.write("### 邮件正文\n\n")
            f.write("----- EMAIL BODY START -----\n")
            f.write(body)
            if body and not body.endswith("\n"):
                f.write("\n")
            f.write("----- EMAIL BODY END -----\n\n")
            f.write("---\n\n")

    return len(ordered)


def export_all_flagged_mail(wb, inbox_records):
    """
    每次运行统一入口：
    1) 从本轮扫描到的邮件里筛选“当前已经 Flag”的全部邮件；
    2) 写入主报告 Flag邮件 Sheet；
    3) 导出完整 JSONL + Markdown；
    4) 在汇总页留下文件位置和数量。
    """
    flagged_records = [r for r in inbox_records if r.get("is_flagged")]
    write_flagged_mail_sheet(wb, flagged_records)
    flagged_count = export_flagged_mail_files(flagged_records)

    ws = wb["汇总"]
    ws["A18"] = "Flag 邮件结构化导出"
    ws["A18"].font = Font(bold=True)
    ws["A19"] = "当前 Flag 邮件数"
    ws["B19"] = flagged_count
    ws["A20"] = "GPT 完整数据（JSONL）"
    ws["B20"] = str(FLAGGED_JSONL_PATH)
    ws["A21"] = "GPT/人眼完整数据（Markdown）"
    ws["B21"] = str(FLAGGED_MD_PATH)
    ws["A22"] = "说明"
    ws["B22"] = "Flag邮件 Sheet 便于筛选；完整正文以 JSONL / Markdown 为准，不受 Excel 32767 字符上限影响。"

    print(
        f"\n[Flag导出] 当前 Flag 邮件 {flagged_count} 封 | "
        f"JSONL: {FLAGGED_JSONL_PATH.name} | MD: {FLAGGED_MD_PATH.name}"
    )
    return flagged_count


def build_domain_sheet(
    wb,
    domains,
    domain_to_emails,
    sent_index,
    inbox_records,
):
    if "域名追踪" in wb.sheetnames:
        del wb["域名追踪"]

    ws = wb.create_sheet("域名追踪")
    ws.append(DOMAIN_HEADERS)

    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DOMAIN_HEADERS))}1"

    domain_start = time.time()
    domain_total = len(domains)
    domain_stats = []
    human_history_rows = []
    inbox_match_index = build_inbox_match_index(inbox_records)

    for domain_index, domain in enumerate(domains, start=1):
        d, known_emails, known_email_domains = build_domain_aliases(
            domain,
            domain_to_emails,
        )

        sent_matches = sent_matches_domain(
            d,
            known_emails,
            known_email_domains,
            sent_index,
        )

        # 严格匹配：保留“为什么这封邮件属于这个域名”的依据。
        matched_pairs = []
        inbox_candidates = candidate_inbox_records_for_domain(
            d, known_emails, sent_matches, inbox_match_index
        )
        for r in inbox_candidates:
            reason = inbox_record_domain_match_reason(
                r,
                d,
                known_emails,
                known_email_domains,
                sent_matches,
                sent_index,
            )
            if reason:
                matched_pairs.append((r, reason))

        bounces = []
        autos = []
        humans = []
        for r, reason in matched_pairs:
            rr = dict(r)
            rr["_domain_match_reason"] = reason
            if rr["category"] == "发送失败":
                bounces.append(rr)
            elif rr["category"] == "自动回复":
                autos.append(rr)
            elif rr["category"] == "真人回复":
                humans.append(rr)
                human_history_rows.append({"domain": d, "record": rr, "reason": reason})

        # 只列出真正属于这个域名的主动触达收件人；不把同封邮件中其它 CC/To 混进来。
        sent_recipient_emails = []
        for rec in sent_matches:
            for e in rec["recipients"]:
                if _email_belongs_to_domain(e, d, known_emails) and e not in sent_recipient_emails:
                    sent_recipient_emails.append(e)

        # 失败邮箱必须再次经过该域名校验。
        failed_emails = []
        for r in bounces:
            failed_candidates = r.get("failed_recipients") or extract_emails(r.get("failed_recipient", ""))
            for e in failed_candidates:
                if _email_belongs_to_domain(e, d, known_emails) and e not in failed_emails:
                    failed_emails.append(e)

        latest_sent = sent_matches[-1]["sent_time"] if sent_matches else ""
        latest_human = max(
            humans,
            key=lambda r: parse_dt(r["received_time"]),
            default=None,
        )

        has_known_email = bool(known_emails or sent_recipient_emails)
        touched = bool(sent_matches)

        if humans:
            latest_outcome = infer_latest_outcome(latest_human)
            final_result = (
                f"{latest_outcome}\n"
                f"最新回复：{latest_human['received_time']} | "
                f"{latest_human['sender_name']} <{latest_human['sender_email']}>\n"
                f"匹配依据：{latest_human.get('_domain_match_reason', '')}\n"
                f"{latest_human['body'][:1200]}"
            )
            next_action = "优先人工阅读真人回复；需要时继续推进合作。"
            fill = GREEN

        elif bounces and touched:
            final_result = "已主动触达，但检测到属于该域名的发送失败/退信。"
            next_action = "重新寻找正确可用邮箱；找不到则改走官网 Contact/Partner 渠道。"
            fill = RED

        elif touched and autos:
            final_result = "已主动触达，仅找到自动回复，未发现真人回复。"
            next_action = "考虑换联系人/换邮箱，或继续通过官网合作入口跟进。"
            fill = YELLOW

        elif touched:
            final_result = "已主动触达，未发现真人回复。"
            next_action = "石沉大海：考虑换联系人、换邮箱或二次跟进。"
            fill = YELLOW

        elif has_known_email:
            final_result = "已找到邮箱，但 Sent Items 未发现主动触达记录。"
            next_action = "可能遗漏发送；核对后优先补发。"
            fill = BLUE

        else:
            final_result = "未找到可用邮箱，也未发现主动触达记录。"
            next_action = "重新寻找邮箱；若仍找不到，改走官网个性化联系方式。"
            fill = PURPLE

        row = [
            domain,
            "; ".join(known_emails),
            "是" if has_known_email else "否",
            "是" if touched else "否",
            len(sent_matches),
            latest_sent,
            "; ".join(sent_recipient_emails),
            "是" if bounces else "否",
            "; ".join(failed_emails),
            len(autos),
            len(humans),
            latest_human["received_time"] if latest_human else "",
            (
                f"{latest_human['sender_name']} <{latest_human['sender_email']}>"
                if latest_human else ""
            ),
            final_result,
            make_history_text(humans),
            next_action,
        ]

        ws.append(row)
        row_num = ws.max_row

        for c in ws[row_num]:
            c.fill = fill
            c.alignment = Alignment(vertical="top", wrap_text=True)

        domain_stats.append({
            "domain": d,
            "has_known_email": has_known_email,
            "touched": touched,
            "sent_count": len(sent_matches),
            "auto_count": len(autos),
            "human_count": len(humans),
            "bounce_count": len(bounces),
            "status": _domain_status(touched, humans),
        })

        if domain_index % 10 == 0 or domain_index == domain_total:
            elapsed = time.time() - domain_start
            speed = domain_index / elapsed if elapsed > 0 else 0
            remaining = domain_total - domain_index
            eta = remaining / speed if speed > 0 else 0

            print(
                f"[Domain] {domain_index}/{domain_total} "
                f"({domain_index / domain_total * 100:.1f}%) | "
                f"当前: {domain} | "
                f"ETA {eta / 60:.1f} min",
                flush=True,
            )

    widths = [34, 45, 14, 14, 15, 22, 45, 14, 40, 14, 14, 22, 38, 85, 110, 70]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 新增：域名级统计、饼图/柱状图，以及完整真人回复历史。
    write_domain_statistics(wb, domain_stats)
    write_human_history_sheet(wb, human_history_rows)

    print(
        f"\n[Domain统计] 总域名 {len(domain_stats)} | "
        f"未触达 {sum(1 for x in domain_stats if not x['touched'])} | "
        f"已触达无真人 {sum(1 for x in domain_stats if x['touched'] and x['human_count'] == 0)} | "
        f"已触达有真人 {sum(1 for x in domain_stats if x['touched'] and x['human_count'] > 0)}"
    )


# ============================================================
# 域名输入
# ============================================================

def load_domains_interactive():
    answer = input(
        "\n是否要查询一列域名的 Outlook 触达/回复情况？ [y/N]: "
    ).strip().lower()

    if answer not in {"y", "yes", "1"}:
        return []

    print("\n域名输入方式：")
    print(f"  1 = 直接粘贴一列纯域名")
    print(f"  2 = 读取 {DOMAINS_PATH}")
    mode = input("请选择 [1/2]: ").strip()

    if mode == "2":
        if not DOMAINS_PATH.exists():
            print(f"找不到：{DOMAINS_PATH}")
            return []
        lines = DOMAINS_PATH.read_text(encoding="utf-8-sig", errors="ignore").splitlines()
        domains = [normalize_domain(x) for x in lines if normalize_domain(x)]
        print(f"已读取 {len(domains)} 个域名。")
        return domains

    print("\n请粘贴纯域名，一行一个；粘贴完成后输入 END 并回车：")
    domains = []
    while True:
        line = input().strip()
        if line.upper() == "END":
            break
        d = normalize_domain(line)
        if d:
            domains.append(d)

    print(f"已读取 {len(domains)} 个域名。")
    return domains


# ============================================================
# 主程序
# ============================================================

def main():
    if sys.platform != "win32":
        raise RuntimeError("需要 Windows + Classic Outlook。")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 72)
    print("ArkSwift Outlook Inbox Triage + Domain Tracking")
    print("=" * 72)
    print("Outlook 修改：仅真人回复 Flag")
    print("不会删除 / 移动 / 标已读 / 自动回复邮件")
    print(f"输出：{REPORT_PATH}")

    # 先询问域名模式
    domains = load_domains_interactive()

    pythoncom.CoInitialize()

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        namespace = outlook.GetNamespace("MAPI")

        account, store = get_account_and_store(namespace)

        try:
            own_email = clean_text(account.SmtpAddress).lower() if account else ""
        except Exception:
            own_email = ""

        print(f"\n邮箱：{own_email or 'Outlook 默认 Store'}")

        inbox = store.GetDefaultFolder(OL_FOLDER_INBOX)
        sent = store.GetDefaultFolder(OL_FOLDER_SENT_MAIL)

        sent_index = build_sent_index(sent)
        domain_to_emails, email_to_domains = load_outreach_mapping()

        if OUTREACH_PATH.exists():
            print(
                f"已读取 outreach_emails.xlsx："
                f"{len(domain_to_emails)} 个 Domain 映射。\n"
            )
        else:
            print("未找到 outreach_emails.xlsx，将只靠 Outlook 记录匹配。\n")

        wb = make_workbook()

        counts = {
            "发送失败": 0,
            "自动回复": 0,
            "真人回复": 0,
            "其他": 0,
        }

        inbox_records = []
        total_scanned = 0
        scan_start = time.time()
        total_unread = 0
        flag_new = 0
        flag_existing = 0
        flag_failed = 0

        folders = list(walk_folder(inbox, SCAN_INBOX_SUBFOLDERS))
        print(f"准备扫描 Inbox 文件夹：{len(folders)} 个\n")

        stop = False

        for folder in folders:
            if stop:
                break

            folder_path = (
                clean_text(safe_get(folder, "FolderPath"))
                or clean_text(safe_get(folder, "Name"))
            )
            print(f"扫描：{folder_path}")

            items = folder.Items
            try:
                items.Sort("[ReceivedTime]", True)
            except Exception:
                pass

            total = items.Count
            limit = total if MAX_INBOX_ITEMS is None else min(total, MAX_INBOX_ITEMS)

            for i in range(1, limit + 1):
                if MAX_INBOX_ITEMS is not None and total_scanned >= MAX_INBOX_ITEMS:
                    stop = True
                    break

                try:
                    item = items.Item(i)
                except Exception:
                    continue

                mc = clean_text(safe_get(item, "MessageClass"))
                if not (mc.startswith("IPM.Note") or mc.startswith("REPORT.")):
                    continue

                subject = clean_text(safe_get(item, "Subject"))
                body = clean_text(safe_get(item, "Body"))
                classification = classify_item(item, sent_index, own_email)

                unread = bool(safe_get(item, "UnRead", False))
                if unread:
                    total_unread += 1

                # 记录“当前是否已经 Flag”。如果本轮新打 Flag，下面会再次读取最终状态。
                was_flagged = is_item_flagged(item)

                flag_result = ""
                if classification["is_human"]:
                    flag_result = mark_human_reply_flag(item)
                    print(
                        f"[FLAG] {flag_result} | "
                        f"{get_sender_email(item)} | "
                        f"{subject[:70]}",
                        flush=True,
                    )
                    if flag_result == "新增Flag":
                        flag_new += 1
                    elif flag_result == "已有Flag":
                        flag_existing += 1
                    elif flag_result.startswith("Flag失败"):
                        flag_failed += 1

                record = {
                    "category": classification["category"],
                    "intent": classification["intent"],
                    "confidence": classification["confidence"],
                    "received_time": fmt_dt(safe_get(item, "ReceivedTime", None)),
                    "unread": "是" if unread else "否",
                    "sender_name": clean_text(safe_get(item, "SenderName")),
                    "sender_email": get_sender_email(item),
                    "subject": subject,
                    "failed_recipient": classification["failed_recipient"],
                    "failed_recipients": classification.get("failed_recipients", []),
                    "reason": classification["reason"],
                    "body_preview": body[:BODY_PREVIEW_CHARS],
                    "body": body,
                    "conversation_id": clean_text(safe_get(item, "ConversationID")),
                    "conversation_topic": clean_text(safe_get(item, "ConversationTopic")),
                    "message_class": mc,
                    "entry_id": clean_text(safe_get(item, "EntryID")),
                    "folder": folder_path,
                    "flag_result": flag_result,
                    # 最终状态：历史已有 Flag 或本轮刚刚新增 Flag 都会是 True。
                    "is_flagged": bool(was_flagged or is_item_flagged(item)),
                }

                inbox_records.append(record)
                wb["全部"].append(triage_record_row(record))
                wb[record["category"]].append(triage_record_row(record))

                counts[record["category"]] += 1
                total_scanned += 1

                if total_scanned % 25 == 0:
                    elapsed = time.time() - scan_start
                    speed = total_scanned / elapsed if elapsed > 0 else 0
                    remaining = max(0, total - total_scanned)
                    eta = remaining / speed if speed > 0 else 0

                    print(
                        f"[Inbox] {total_scanned}/{total} "
                        f"({total_scanned / total * 100:.1f}%) | "
                        f"失败 {counts['发送失败']} | "
                        f"自动 {counts['自动回复']} | "
                        f"真人 {counts['真人回复']} | "
                        f"其他 {counts['其他']} | "
                        f"Flag新增 {flag_new} | "
                        f"ETA {eta / 60:.1f} min",
                        flush=True,
                    )

        write_summary(
            wb,
            counts,
            total_unread,
            flag_new,
            flag_existing,
            flag_failed,
        )

        for name in ["全部", "发送失败", "自动回复", "真人回复", "其他"]:
            apply_triage_widths(wb[name])

        # 每次执行都导出 Outlook 当前全部 Flag 邮件，与是否输入域名无关。
        flagged_count = export_all_flagged_mail(wb, inbox_records)

        if domains:
            print(f"\n正在生成 {len(domains)} 行域名追踪表...")
            build_domain_sheet(
                wb,
                domains,
                domain_to_emails,
                sent_index,
                inbox_records,
            )

        try:
            wb.save(REPORT_PATH)
        except PermissionError:
            raise RuntimeError(
                f"无法写入 {REPORT_PATH}。\n"
                "请先关闭正在打开的 outlook_triage.xlsx。"
            )

        print("\n" + "=" * 72)
        print("整理完成")
        print("=" * 72)
        print(f"共扫描：{total_scanned}")
        print(f"其中未读：{total_unread}")
        print(f"发送失败：{counts['发送失败']}")
        print(f"自动回复：{counts['自动回复']}")
        print(f"真人回复：{counts['真人回复']}")
        print(f"其他：{counts['其他']}")
        print(f"新增真人 Flag：{flag_new}")
        print(f"已有 Flag：{flag_existing}")
        print(f"Flag 失败：{flag_failed}")
        print(f"当前 Flag 邮件导出：{flagged_count}")
        print(f"Flag JSONL：{FLAGGED_JSONL_PATH}")
        print(f"Flag Markdown：{FLAGGED_MD_PATH}")
        if domains:
            print(f"域名追踪：{len(domains)} 行")
        print(f"\n报告：{REPORT_PATH}")

    finally:
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n用户停止。")
    except Exception as exc:
        print("\n脚本停止：")
        print(exc)
        sys.exit(1)
